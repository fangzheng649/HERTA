"""
文档解析模块

支持解析协议文档：
- .docx 使用 python-docx
- .pdf 使用 PyMuPDF
- .doc 不支持，建议转换为 docx/pdf
"""

from __future__ import annotations

from pathlib import Path
from typing import List, Optional, Dict, Any, TYPE_CHECKING

from herta.utils.text import normalize_text, strip_extra_blank_lines
from herta.logging import get_logger

if TYPE_CHECKING:
    from herta.kb.schemas import SourceDocument

logger = get_logger(__name__)

# 尝试导入文档解析库
try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    logger.warning("python-docx 未安装，docx 解析功能不可用")

try:
    import fitz  # PyMuPDF
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False
    logger.warning("PyMuPDF 未安装，pdf 解析功能不可用")


class DocParser:
    """
    文档解析器

    支持 docx 和 pdf 格式
    """

    SUPPORTED_EXTS = ['.docx', '.pdf']

    def __init__(self):
        self.stats = {
            "parsed": 0,
            "failed": 0,
            "unsupported": 0,
        }

    def parse(self, file_path: Path, source: str) -> List[SourceDocument]:
        """
        解析文档

        Args:
            file_path: 文档路径
            source: 来源标识

        Returns:
            SourceDocument 列表

        Raises:
            ValueError: 不支持的文件格式
            FileNotFoundError: 文件不存在
        """
        file_path = Path(file_path)

        if not file_path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")

        ext = file_path.suffix.lower()

        if ext == '.docx':
            return self._parse_docx(file_path, source)
        elif ext == '.pdf':
            return self._parse_pdf(file_path, source)
        elif ext == '.doc':
            raise ValueError(
                f"不支持 .doc 格式: {file_path}\n"
                f"当前版本建议先将 .doc 转换为 .docx 或 .pdf 格式后再解析。\n"
                f"转换方式：使用 Microsoft Word 或 LibreOffice 另存为 docx/pdf"
            )
        else:
            raise ValueError(
                f"不支持的文件格式: {ext}\n"
                f"支持格式: {', '.join(self.SUPPORTED_EXTS)}"
            )

    def _parse_docx(self, file_path: Path, source: str) -> List[SourceDocument]:
        """
        解析 docx 文件

        按段落解析，保留段落结构
        """
        from herta.kb.schemas import SourceDocument

        if not DOCX_AVAILABLE:
            raise ImportError("python-docx 未安装，无法解析 docx 文件")

        logger.info(f"解析 docx: {file_path}")

        try:
            doc = Document(file_path)
        except Exception as e:
            logger.error(f"打开 docx 文件失败: {file_path}, 错误: {e}")
            raise

        documents = []
        para_index = 0

        # 遍历所有段落
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            # 清洗文本
            text = normalize_text(text, preserve_newlines=True)

            if text:
                doc_obj = SourceDocument(
                    source=source,
                    file=str(file_path),
                    kind="doc",
                    section=f"paragraph_{para_index:04d}",
                    text=text,
                    metadata={
                        "style": para.style.name if para.style else "Normal",
                        "index": para_index,
                    }
                )
                documents.append(doc_obj)
                para_index += 1

        # 同时提取表格内容
        for table_idx, table in enumerate(doc.tables):
            table_texts = []
            for row in table.rows:
                row_text = " | ".join(
                    normalize_text(cell.text.strip())
                    for cell in row.cells
                )
                if row_text.strip():
                    table_texts.append(row_text)

            if table_texts:
                table_content = "\n".join(table_texts)
                doc_obj = SourceDocument(
                    source=source,
                    file=str(file_path),
                    kind="doc",
                    section=f"table_{table_idx:04d}",
                    text=table_content,
                    metadata={
                        "type": "table",
                        "table_index": table_idx,
                        "row_count": len(table.rows),
                    }
                )
                documents.append(doc_obj)

        self.stats["parsed"] += 1
        logger.info(f"docx 解析完成: {file_path} -> {len(documents)} 个段落/表格")
        return documents

    def _parse_pdf(self, file_path: Path, source: str) -> List[SourceDocument]:
        """
        解析 pdf 文件

        按页解析，每页作为一个 document
        """
        from herta.kb.schemas import SourceDocument

        if not PDF_AVAILABLE:
            raise ImportError("PyMuPDF 未安装，无法解析 pdf 文件")

        logger.info(f"解析 pdf: {file_path}")

        documents = []

        try:
            with fitz.open(file_path) as pdf_doc:
                for page_num in range(len(pdf_doc)):
                    page = pdf_doc[page_num]

                    # 提取文本
                    text = page.get_text()
                    text = normalize_text(text, preserve_newlines=True)
                    text = strip_extra_blank_lines(text)

                    if text.strip():
                        doc_obj = SourceDocument(
                            source=source,
                            file=str(file_path),
                            kind="doc",
                            section=f"page_{page_num:04d}",
                            text=text,
                            metadata={
                                "page_number": page_num + 1,
                                "page_count": len(pdf_doc),
                            }
                        )
                        documents.append(doc_obj)

        except Exception as e:
            logger.error(f"解析 pdf 失败: {file_path}, 错误: {e}")
            self.stats["failed"] += 1
            raise

        self.stats["parsed"] += 1
        logger.info(f"pdf 解析完成: {file_path} -> {len(documents)} 页")
        return documents

    def can_parse(self, file_path: Path) -> bool:
        """
        检查是否可以解析该文件

        Args:
            file_path: 文件路径

        Returns:
            是否支持解析
        """
        ext = Path(file_path).suffix.lower()
        if ext == '.docx' and not DOCX_AVAILABLE:
            return False
        if ext == '.pdf' and not PDF_AVAILABLE:
            return False
        return ext in self.SUPPORTED_EXTS


def parse_document(file_path: str, source: str) -> List[SourceDocument]:
    """
    便捷函数：解析单个文档

    Args:
        file_path: 文档路径
        source: 来源标识

    Returns:
        SourceDocument 列表

    Example:
        >>> docs = parse_document("data/source_proto/gdwx/gdwx.docx", "gdwx")
        >>> print(len(docs))
    """
    parser = DocParser()
    return parser.parse(Path(file_path), source)
